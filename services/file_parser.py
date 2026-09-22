import io
import logging
from pathlib import Path
from typing import Union

import docx
from pypdf import PdfReader

logger = logging.getLogger(__name__)


def extract_text_from_pdf(file_path: Path) -> str:
    """
    Extract readable text from a PDF document.
    """
    text_chunks = []
    try:
        reader = PdfReader(str(file_path))
        for page_idx, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text_chunks.append(page_text)
    except Exception as exc:
        logger.error("Error extracting text from PDF (%s): %s", file_path, exc)
        raise ValueError(f"Failed to read PDF file: {exc}") from exc

    return "\n".join(text_chunks).strip()


def extract_text_from_docx(file_path: Path) -> str:
    """
    Extract readable text from a DOCX document.
    """
    text_chunks = []
    try:
        doc = docx.Document(str(file_path))
        # Extract paragraph text
        for para in doc.paragraphs:
            if para.text.strip():
                text_chunks.append(para.text.strip())

        # Extract text from tables if present
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_chunks.append(" | ".join(row_text))
    except Exception as exc:
        logger.error("Error extracting text from DOCX (%s): %s", file_path, exc)
        raise ValueError(f"Failed to read DOCX file: {exc}") from exc

    return "\n".join(text_chunks).strip()


def extract_text_from_txt(file_path: Path) -> str:
    """
    Extract text from a plain text file.
    """
    try:
        return file_path.read_text(encoding="utf-8", errors="ignore").strip()
    except Exception as exc:
        logger.error("Error extracting text from TXT (%s): %s", file_path, exc)
        raise ValueError(f"Failed to read text file: {exc}") from exc


def extract_text_from_file(file_path: Union[str, Path]) -> str:
    """
    Dispatcher to extract raw text based on file format.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    suffix = path.suffix.lower()

    if suffix == ".pdf":
        text = extract_text_from_pdf(path)
    elif suffix in {".docx", ".doc"}:
        text = extract_text_from_docx(path)
    elif suffix in {".txt", ".text"}:
        text = extract_text_from_txt(path)
    else:
        raise ValueError(f"Unsupported file format: {suffix}")

    return text